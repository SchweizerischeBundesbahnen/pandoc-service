"""Visual tests: a PDF made from a DOCX keeps the page of the DOCX.

Each test builds a DOCX, converts it to PDF through the service, and compares
the pages with the references (see tests/visual.py).
"""

from __future__ import annotations

import io

from docx import Document
from docx.shared import Cm
from PIL import Image

from tests.test_container import TestParameters
from tests.visual import assert_pages_match

DOCX_MIME = "application/vnd.openxmlformats-officedocument.wordprocessingml.document"


def _png(width: int, height: int, color: tuple[int, int, int]) -> io.BytesIO:
    """A PNG of one color with a black frame, so its edges show on the page."""
    image = Image.new("RGB", (width, height), (0, 0, 0))
    image.paste(color, (4, 4, width - 4, height - 4))
    buffer = io.BytesIO()
    image.save(buffer, format="PNG")
    buffer.seek(0)
    return buffer


def _docx_to_pdf(test_parameters: TestParameters, document) -> bytes:
    docx = io.BytesIO()
    document.save(docx)
    response = test_parameters.request_session.post(f"{test_parameters.base_url}/convert/docx/to/pdf", files={"source": ("in.docx", docx.getvalue(), DOCX_MIME)})
    assert response.status_code == 200, response.text
    return response.content


def test_a5_page_with_narrow_margins(test_parameters: TestParameters):
    """An A5 page with 2 cm margins: the image spans the text width, margin to margin."""
    document = Document()
    section = document.sections[0]
    section.page_width, section.page_height = Cm(14.8), Cm(21)
    section.top_margin = section.bottom_margin = section.left_margin = section.right_margin = Cm(2)
    document.add_paragraph("The image below is as wide as the text.")
    document.add_picture(_png(400, 200, (40, 90, 170)), width=Cm(10.8))
    document.add_paragraph("After the image.")

    assert_pages_match("page_geometry_a5_narrow_margins", _docx_to_pdf(test_parameters, document))
