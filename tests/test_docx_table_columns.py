"""Unit tests for :mod:`app.docx_table_columns`."""

import io
import struct
import zlib

import pytest
from docx import Document
from docx.oxml import parse_xml
from docx.oxml.ns import nsdecls
from docx.shared import Inches

from app.docx_table_columns import decide, distribute

_W = "http://schemas.openxmlformats.org/wordprocessingml/2006/main"
INCH = 914400
MARGINS = 2 * 108 * 635


def _margins(_tc) -> int:
    return MARGINS


def _png() -> bytes:
    def chunk(tag: bytes, payload: bytes) -> bytes:
        return struct.pack(">I", len(payload)) + tag + payload + struct.pack(">I", zlib.crc32(tag + payload))

    raw = b"".join(b"\x00" + b"\xff\xff\xff" * 4 for _ in range(2))
    return b"\x89PNG\r\n\x1a\n" + chunk(b"IHDR", struct.pack(">IIBBBBB", 4, 2, 8, 2, 0, 0, 0)) + chunk(b"IDAT", zlib.compress(raw)) + chunk(b"IEND", b"")


def _table(columns: int, rows: int = 1):
    doc = Document()
    return doc, doc.add_table(rows=rows, cols=columns)


# ----------------------- distribute -----------------------


def test_with_room_for_every_column_the_rest_goes_to_those_without_a_width():
    assert distribute([10, 10], [100, 300], 600, [True, False]) == [100, 500]


def test_with_room_for_every_column_and_all_stated_the_rest_is_shared_by_their_widest():
    assert distribute([10, 10], [100, 300], 800, [True, True]) == [200, 600]


def test_short_of_room_each_column_grows_from_its_narrowest_by_what_it_lacks():
    # 200 to share; the first would grow by 100, the second by 300
    assert distribute([100, 100], [200, 400], 400, [False, False]) == [150, 250]


def test_without_room_even_for_the_narrowest_all_are_scaled_down_alike():
    assert distribute([300, 100], [600, 200], 200, [False, False]) == [150, 50]


def test_empty_columns_share_the_width_evenly():
    assert distribute([0, 0], [0, 0], 100, [False, False]) == [50, 50]


# ----------------------- decide -----------------------


def test_widths_stated_for_every_column_are_taken_as_they_are():
    _, table = _table(2)

    assert decide(table._tbl, 10000, [("pct", 1000), ("pct", 4000)], _margins) == [2000, 8000]


def test_absolute_stated_widths_are_scaled_to_the_table():
    _, table = _table(2)

    assert decide(table._tbl, 6000, [("dxa", 1000), ("dxa", 2000)], _margins) == [2000, 4000]


def test_a_column_holding_an_image_takes_the_room_its_neighbour_does_not_need():
    """The even share of pandoc's grid would leave the image half the table."""
    _, table = _table(2)
    table.cell(0, 0).paragraphs[0].add_run("Short")
    table.cell(0, 1).paragraphs[0].add_run().add_picture(io.BytesIO(_png()), width=Inches(9))

    widths = decide(table._tbl, int(6.5 * INCH), None, _margins)

    assert widths is not None
    assert sum(widths) == pytest.approx(6.5 * INCH, abs=2)
    assert widths[1] > 0.8 * 6.5 * INCH
    assert widths[0] >= MARGINS


def test_a_stated_column_keeps_its_width_and_the_others_share_the_rest():
    _, table = _table(2)
    table.cell(0, 1).paragraphs[0].add_run().add_picture(io.BytesIO(_png()), width=Inches(9))

    widths = decide(table._tbl, 10 * INCH, [("pct", 1000), None], _margins)

    assert widths is not None
    assert widths[0] == 2 * INCH


def test_stated_widths_which_do_not_match_the_grid_are_ignored():
    _, table = _table(2)
    table.cell(0, 0).paragraphs[0].add_run("Short")
    table.cell(0, 1).paragraphs[0].add_run().add_picture(io.BytesIO(_png()), width=Inches(9))

    widths = decide(table._tbl, int(6.5 * INCH), [("pct", 5000)], _margins)

    assert widths is not None
    assert widths[1] > widths[0]


def test_an_image_in_a_crowded_table_keeps_its_width():
    _, table = _table(3)
    for index in range(2):
        table.cell(0, index).paragraphs[0].add_run(" ".join(["word"] * 400))
    table.cell(0, 2).paragraphs[0].add_run().add_picture(io.BytesIO(_png()), width=Inches(5))

    widths = decide(table._tbl, int(6.5 * INCH), None, _margins)

    assert widths is not None
    assert widths[2] >= 5 * INCH + MARGINS


def test_an_image_keeps_its_width_beside_a_long_line_of_text():
    """A browser keeps an image of a stated width and wraps the text beside it; #277."""
    _, table = _table(2)
    table.cell(0, 0).paragraphs[0].add_run().add_picture(io.BytesIO(_png()), width=Inches(160 / 96))
    table.cell(0, 1).paragraphs[0].add_run("The image keeps its 160 pixels, the description takes the rest of the width. " * 3)

    widths = decide(table._tbl, int(6.27 * INCH), None, _margins)

    assert widths is not None
    assert widths[0] >= int(160 / 96 * INCH) + MARGINS


def test_an_image_wider_than_the_table_leaves_the_text_its_longest_word():
    _, table = _table(2)
    table.cell(0, 0).paragraphs[0].add_run("Description")
    table.cell(0, 1).paragraphs[0].add_run().add_picture(io.BytesIO(_png()), width=Inches(9))
    widths = decide(table._tbl, int(6.5 * INCH), None, _margins)

    assert widths is not None
    assert sum(widths) == pytest.approx(int(6.5 * INCH), abs=2)
    # "Description" at the default 10 pt is some 0.8 inch, which its column keeps whole
    assert widths[0] >= 11 * 10 * 12700 // 2 + MARGINS


@pytest.mark.parametrize("stated", [[("pct", 2500), ("pct", 2500), None], [("dxa", 6000), ("dxa", 6000), None]])
def test_a_picture_beside_stated_columns_which_fill_the_table_keeps_a_share(stated):
    """Stated widths of 100 % or more leave no room; the open column still gets a share, as before #277."""
    _, table = _table(3)
    table.cell(0, 2).paragraphs[0].add_run().add_picture(io.BytesIO(_png()), width=Inches(2))

    widths = decide(table._tbl, int(6.5 * INCH), stated, _margins)

    assert widths is not None
    assert sum(widths) == pytest.approx(int(6.5 * INCH), abs=3)
    # Its inch, scaled down alike with the stated columns to the width of the table
    stated_width = 0.5 * 6.5 * INCH if stated[0][0] == "pct" else 6000 * 635  # 2500 fiftieths of a percent are 50 %
    assert widths[2] >= INCH * widths[0] / stated_width - 2


def test_a_text_column_beside_a_stated_column_of_the_full_width_keeps_a_share():
    _, table = _table(2)
    table.cell(0, 1).paragraphs[0].add_run("Description of some length")

    widths = decide(table._tbl, int(6.5 * INCH), [("pct", 5000), None], _margins)

    assert widths is not None
    assert widths[1] > 0.3 * INCH


def test_a_long_word_beside_a_picture_wider_than_the_table_stays_whole():
    """The picture gives up the room the table lacks before a word breaks."""
    _, table = _table(2)
    word = "x" * 58  # 58 characters of 5 pt at the default 10 pt: some 4 inches
    table.cell(0, 0).paragraphs[0].add_run(word)
    table.cell(0, 1).paragraphs[0].add_run().add_picture(io.BytesIO(_png()), width=Inches(9))

    widths = decide(table._tbl, int(6.5 * INCH), None, _margins)

    assert widths is not None
    assert widths[0] >= 58 * 5 * 12700 + MARGINS
    assert widths[1] >= INCH


def test_a_merged_row_widens_the_floor_of_the_column_which_has_room_for_it():
    """A cell spanning a picture and a word raises the floor where the picture has width to give, not evenly."""
    _, table = _table(2, rows=2)
    table.cell(0, 0).paragraphs[0].add_run().add_picture(io.BytesIO(_png()), width=Inches(6))
    table.cell(0, 1).paragraphs[0].add_run("x" * 14)  # some 1 inch
    table.cell(1, 0).merge(table.cell(1, 1)).paragraphs[0].add_run("y" * 58)  # some 4 inches

    widths = decide(table._tbl, int(6.5 * INCH), None, lambda _tc: 0)

    assert widths is not None
    # The word keeps its inch, and the picture gives up only the half inch the table lacks
    assert widths[0] >= int(5.4 * INCH)


def test_a_word_split_into_runs_by_its_formatting_is_one_word():
    _, table = _table(2)
    paragraph = table.cell(0, 0).paragraphs[0]
    paragraph.add_run("x" * 29)
    paragraph.add_run("x" * 29).bold = True  # one word of 58 characters, some 4 inches
    table.cell(0, 1).paragraphs[0].add_run().add_picture(io.BytesIO(_png()), width=Inches(9))

    widths = decide(table._tbl, int(6.5 * INCH), None, _margins)

    assert widths is not None
    assert widths[0] >= 58 * 5 * 12700 + MARGINS


def test_a_no_break_space_does_not_end_a_word():
    _, table = _table(2)
    table.cell(0, 0).paragraphs[0].add_run("x" * 29 + "\u00a0" + "x" * 28)  # 58 characters on one line
    table.cell(0, 1).paragraphs[0].add_run().add_picture(io.BytesIO(_png()), width=Inches(9))

    widths = decide(table._tbl, int(6.5 * INCH), None, _margins)

    assert widths is not None
    assert widths[0] >= 58 * 5 * 12700 + MARGINS


def test_a_space_between_runs_ends_a_word():
    _, table = _table(2)
    paragraph = table.cell(0, 0).paragraphs[0]
    paragraph.add_run("x" * 29 + " ")
    paragraph.add_run("x" * 29)
    table.cell(0, 1).paragraphs[0].add_run().add_picture(io.BytesIO(_png()), width=Inches(9))

    widths = decide(table._tbl, int(6.5 * INCH), None, _margins)

    assert widths is not None
    assert widths[0] < 58 * 5 * 12700


def test_overlapping_merged_cells_share_the_column_between_them():
    """Two merged words overlapping on the middle column both keep their room beside two pictures."""
    _, table = _table(3, rows=3)
    for column in (0, 2):
        table.cell(0, column).paragraphs[0].add_run().add_picture(io.BytesIO(_png()), width=Inches(6))
    table.cell(1, 0).merge(table.cell(1, 1)).paragraphs[0].add_run("y" * 58)  # some 4 inches
    table.cell(2, 1).merge(table.cell(2, 2)).paragraphs[0].add_run("z" * 58)  # some 4 inches

    widths = decide(table._tbl, int(6.5 * INCH), None, lambda _tc: 0)

    assert widths is not None
    word = 58 * 5 * 12700
    assert widths[0] + widths[1] >= word
    assert widths[1] + widths[2] >= word
    assert sum(widths) == pytest.approx(int(6.5 * INCH), abs=3)


def test_a_cell_spanning_columns_widens_them_by_what_they_lack():
    _, table = _table(2, rows=2)
    merged = table.cell(0, 0).merge(table.cell(0, 1))
    merged.paragraphs[0].add_run().add_picture(io.BytesIO(_png()), width=Inches(4))

    widths = decide(table._tbl, 10 * INCH, None, _margins)

    assert widths is not None
    assert sum(widths) == pytest.approx(10 * INCH, abs=2)
    assert widths[0] == pytest.approx(widths[1], abs=2)


def test_a_nested_table_needs_the_room_of_its_columns():
    _, table = _table(2)
    table.cell(0, 0).paragraphs[0].add_run("Short")
    inner = table.cell(0, 1).add_table(rows=1, cols=2)
    for index in range(2):
        inner.cell(0, index).paragraphs[0].add_run().add_picture(io.BytesIO(_png()), width=Inches(2))

    widths = decide(table._tbl, int(6.5 * INCH), None, _margins)

    assert widths is not None
    assert widths[1] >= 4 * INCH + 3 * MARGINS


def test_text_in_a_larger_font_needs_a_wider_column():
    doc, table = _table(2)
    table.cell(0, 0).paragraphs[0].add_run(" ".join(["word"] * 10))
    table.cell(0, 1).paragraphs[0].add_run().add_picture(io.BytesIO(_png()), width=Inches(9))
    styles = doc.styles.element
    small = decide(table._tbl, int(6.5 * INCH), None, _margins, styles)
    defaults = styles.find(f"{{{_W}}}docDefaults")
    if defaults is None:
        defaults = parse_xml(f"<w:docDefaults {nsdecls('w')}/>")
        styles.insert(0, defaults)
    for child in list(defaults):
        defaults.remove(child)
    defaults.append(parse_xml(f'<w:rPrDefault {nsdecls("w")}><w:rPr><w:sz w:val="48"/></w:rPr></w:rPrDefault>'))
    for style in styles.findall(f"{{{_W}}}style"):
        size = style.find(f"{{{_W}}}rPr/{{{_W}}}sz")
        if size is not None:
            size.getparent().remove(size)

    large = decide(table._tbl, int(6.5 * INCH), None, _margins, styles)

    assert small is not None
    assert large is not None
    assert large[0] > small[0]


def test_a_table_without_columns_has_no_widths():
    table = parse_xml(f"<w:tbl {nsdecls('w')}><w:tblPr/></w:tbl>")

    assert decide(table, INCH, None, _margins) is None


def test_a_table_without_a_grid_is_laid_out_with_no_more_columns_than_html_allows():
    from app.docx_table_columns import MAX_COLUMNS

    table = parse_xml(f'<w:tbl {nsdecls("w")}><w:tblPr/><w:tr><w:tc><w:tcPr><w:gridSpan w:val="100000000"/></w:tcPr><w:p/></w:tc></w:tr></w:tbl>')

    widths = decide(table, MAX_COLUMNS * 10, None, _margins)

    assert widths is not None
    assert len(widths) == MAX_COLUMNS


def test_a_template_stating_a_size_or_a_span_int_refuses_does_not_fail_the_layout():
    """A superscript two passes str.isdigit and fails int; the defaults stand instead."""
    table = parse_xml(
        f'<w:tbl {nsdecls("w")}><w:tblPr/><w:tblGrid><w:gridCol w:w="1"/><w:gridCol w:w="1"/></w:tblGrid>'
        f'<w:tr><w:tc><w:tcPr><w:gridSpan w:val="²"/></w:tcPr><w:p><w:r><w:rPr><w:sz w:val="²"/></w:rPr><w:t>Short</w:t></w:r></w:p></w:tc>'
        f"<w:tc><w:p/></w:tc></w:tr></w:tbl>"
    )
    styles = parse_xml(f'<w:styles {nsdecls("w")}><w:docDefaults><w:rPrDefault><w:rPr><w:sz w:val="{"9" * 5000}"/></w:rPr></w:rPrDefault></w:docDefaults></w:styles>')

    widths = decide(table, 10 * INCH, None, _margins, styles)

    assert widths is not None
    assert len(widths) == 2
