"""Tests for app/docx_page_geometry.py: the page of a DOCX as LaTeX geometry."""

from __future__ import annotations

import io
import zipfile

from docx import Document
from docx.shared import Cm, Inches

from app.docx_page_geometry import geometry_variables


def _docx(width=None, height=None, top=None, bottom=None, left=None, right=None) -> bytes:
    document = Document()
    section = document.sections[0]
    for name, value in (("page_width", width), ("page_height", height), ("top_margin", top), ("bottom_margin", bottom), ("left_margin", left), ("right_margin", right)):
        if value is not None:
            setattr(section, name, value)
    document.add_paragraph("text")
    buffer = io.BytesIO()
    document.save(buffer)
    return buffer.getvalue()


def _as_dict(variables: list[str]) -> dict[str, str]:
    assert variables[::2] == ["-V"] * (len(variables) // 2)
    return dict(item.removeprefix("geometry:").split("=") for item in variables[1::2])


def _edit_document(docx: bytes, old: bytes, new: bytes) -> bytes:
    """The same package with one replacement in word/document.xml."""
    source = zipfile.ZipFile(io.BytesIO(docx))
    buffer = io.BytesIO()
    with zipfile.ZipFile(buffer, "w") as target:
        for name in source.namelist():
            data = source.read(name)
            target.writestr(name, data.replace(old, new) if name == "word/document.xml" else data)
    return buffer.getvalue()


def _without_page_geometry(docx: bytes) -> bytes:
    """The same package with <w:pgSz> and <w:pgMar> removed, as pandoc writes it."""
    source = zipfile.ZipFile(io.BytesIO(docx))
    buffer = io.BytesIO()
    with zipfile.ZipFile(buffer, "w") as target:
        for name in source.namelist():
            data = source.read(name)
            if name == "word/document.xml":
                text = data.decode()
                for tag in ("w:pgSz", "w:pgMar"):
                    start = text.index(f"<{tag} ")
                    text = text[:start] + text[text.index("/>", start) + 2 :]
                data = text.encode()
            target.writestr(name, data)
    return buffer.getvalue()


def test_reads_the_page_size_and_margins():
    variables = _as_dict(geometry_variables(_docx(Cm(14.8), Cm(21), Cm(2), Cm(2.5), Cm(3), Cm(1.5))))
    assert variables == {
        "paperwidth": "5.8271in",
        "paperheight": "8.2681in",
        "top": "0.7875in",
        "right": "0.5903in",
        "bottom": "0.9840in",
        "left": "1.1812in",
    }


def test_a_document_without_page_geometry_gets_letter_with_one_inch_margins():
    variables = _as_dict(geometry_variables(_without_page_geometry(_docx())))
    assert variables == {"paperwidth": "8.5000in", "paperheight": "11.0000in", "top": "1.0000in", "right": "1.0000in", "bottom": "1.0000in", "left": "1.0000in"}


def test_the_first_section_decides():
    """A section a paragraph closes comes before the body's own sectPr."""
    document = Document()
    first = document.sections[0]
    first.page_width, first.page_height = Inches(5), Inches(7)
    document.add_paragraph("first section")
    second = document.add_section()
    second.page_width, second.page_height = Inches(11), Inches(8.5)
    buffer = io.BytesIO()
    document.save(buffer)

    variables = _as_dict(geometry_variables(buffer.getvalue()))
    assert (variables["paperwidth"], variables["paperheight"]) == ("5.0000in", "7.0000in")


def test_a_negative_margin_counts_by_its_size():
    docx = _edit_document(_docx(top=Inches(1)), b'w:top="1440"', b'w:top="-1440"')
    assert _as_dict(geometry_variables(docx))["top"] == "1.0000in"


def test_bytes_which_are_no_docx_get_the_fallback():
    assert _as_dict(geometry_variables(b"not a zip"))["paperwidth"] == "8.5000in"


def test_a_package_without_a_document_part_gets_the_fallback():
    buffer = io.BytesIO()
    with zipfile.ZipFile(buffer, "w") as target:
        target.writestr("other.xml", "<x/>")
    assert _as_dict(geometry_variables(buffer.getvalue()))["paperheight"] == "11.0000in"


def test_a_value_which_is_no_number_gets_the_fallback():
    docx = _edit_document(_docx(width=Inches(6)), b'w:w="8640"', b'w:w="wide"')
    assert _as_dict(geometry_variables(docx))["paperwidth"] == "8.5000in"


def test_a_document_part_without_a_body_gets_the_fallback():
    buffer = io.BytesIO()
    with zipfile.ZipFile(buffer, "w") as target:
        target.writestr("word/document.xml", '<w:document xmlns:w="http://schemas.openxmlformats.org/wordprocessingml/2006/main"/>')
    assert _as_dict(geometry_variables(buffer.getvalue()))["paperwidth"] == "8.5000in"
